import openpyxl
import os

file_path = "C:/Users/DanielFelipeGómezBra/Universidad EAFIT/bd - Documentos/bd_eafit.xlsx"

try:
    if not os.path.exists(file_path):
        print(f"Error: El archivo no existe en {file_path}")
    else:
        wb = openpyxl.load_workbook(file_path, data_only=True)
        print(f"Archivo abierto correctamente: {file_path}")
        print("--- Hojas ---")
        print(wb.sheetnames)
        
        power_apps_sources = {
            "Cocina": ["Elemento", "Und", "Existencia", "Bueno", "Regular", "Malo", "Observaciones", "Observaciones2"],
            "Prestamos": ["Item", "Tipo", "Usuario", "Fecha", "FechaFin", "Notas", "Estado"],
            "Recursos": ["Juegos", "Descripción", "Unidades", "Categoria", "Observaciones"],
            "Reservas": ["Docente", "Recurso", "Bloque", "FechaInicio", "FechaFin", "Notas", "Notas2", "Estado", "Aprobado"],
            "Usuarios": ["Nombre", "Email", "Rol", "Activo"]
        }
        
        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            print(f"\n--- Hoja: {sheet_name} ---")
            tables = sheet.tables
            if not tables:
                print("No se encontraron tablas estructuradas en esta hoja.")
            else:
                for table_name, table_range in tables.items():
                    print(f"Tabla: {table_name}, Rango: {table_range}")
                    # Obtener encabezados
                    min_col, min_row, max_col, max_row = openpyxl.utils.range_boundaries(table_range)
                    headers = [sheet.cell(row=min_row, column=c).value for c in range(min_col, max_col + 1)]
                    print(f"Encabezados: {headers}")
                    
                    # Comparar con Power Apps si el nombre coincide o se parece
                    source_key = next((k for k in power_apps_sources.keys() if k.lower() in table_name.lower() or table_name.lower() in k.lower()), None)
                    if source_key:
                        expected = power_apps_sources[source_key]
                        missing = [h for h in expected if h not in headers]
                        extra = [h for h in headers if h not in expected]
                        print(f"Comparación con origen Power Apps '{source_key}':")
                        if not missing and not extra:
                            print(" - Coincidencia exacta.")
                        else:
                            if missing: print(f" - Faltan en Excel: {missing}")
                            if extra: print(f" - Sobran en Excel (con respecto al origen): {extra}")

except Exception as e:
    print(f"Error al procesar el archivo: {e}")
